import datetime,json,pathlib,urllib.request
from openpyxl import Workbook,load_workbook
ROOT=pathlib.Path(__file__).resolve().parents[1]
DIR=ROOT/"league_archive"
URL="https://lm-api-reads.fantasy.espn.com/apis/v3/games/fba/seasons/2027/segments/0/leagues/88948640?view=mTeam&view=mRoster&view=mStatus&view=mMatchup&view=mTransactions2"
def main():
 req=urllib.request.Request(URL,headers={"User-Agent":"Mozilla/5.0","Accept":"application/json"})
 with urllib.request.urlopen(req,timeout=45) as r: data=json.load(r)
 teams=data.get("teams",[])
 if data.get("id")!=88948640 or len(teams)!=16:raise RuntimeError("Unexpected ESPN league or incomplete teams")
 now=datetime.datetime.now(datetime.timezone.utc).isoformat()
 current={}
 for t in teams:
  tid=str(t["id"])
  entries=t.get("roster",{}).get("entries")
  if not entries:raise RuntimeError("Roster missing for "+tid)
  name=t.get("name") or (str(t.get("location",""))+" "+str(t.get("nickname",""))).strip()
  current[tid]={"name":name,"players":[{"id":e.get("playerId"),"name":(e.get("playerPoolEntry",{}).get("player") or {}).get("fullName"),"slot":e.get("lineupSlotId"),"injury":(e.get("playerPoolEntry",{}).get("player") or {}).get("injuryStatus")} for e in entries]}
 oldfile=DIR/"latest.json"
 old=json.loads(oldfile.read_text()) if oldfile.exists() else {}
 changes=[]
 for tid,t in current.items():
  before={str(p["id"]):p for p in old.get("teams",{}).get(tid,{}).get("players",[])}
  after={str(p["id"]):p for p in t["players"]}
  if old:
   for pid in after.keys()-before.keys():changes.append((now,t["name"],"ADD",after[pid]["name"],pid))
   for pid in before.keys()-after.keys():changes.append((now,t["name"],"DROP",before[pid]["name"],pid))
 result={"updatedAt":now,"leagueId":88948640,"teams":current}
 oldfile.write_text(json.dumps(result,ensure_ascii=False,indent=2)+"\n")
 with (DIR/"events.jsonl").open("a",encoding="utf-8") as f:
  for event in changes:f.write(json.dumps(event,ensure_ascii=False)+"\n")
 book=DIR/"ALTINTEPE_WARRIORS_LIVE_LEAGUE_2026-27.xlsx"
 wb=load_workbook(book) if book.exists() else Workbook()
 ws=wb.active
 ws.title="CANLI_KADROLAR"
 ws.delete_rows(1,ws.max_row)
 ws.append(["Kontrol UTC","Takım ID","Takım","Oyuncu ID","Oyuncu","Kadro Pozisyonu","Sakatlık"])
 for tid,t in sorted(current.items(),key=lambda p:int(p[0])):
  for p in t["players"]:ws.append([now,int(tid),t["name"],p["id"],p["name"],p["slot"],p["injury"]])
 hist=wb["HAREKET_GECMISI"] if "HAREKET_GECMISI" in wb else wb.create_sheet("HAREKET_GECMISI")
 if hist.max_row==1 and hist["A1"].value is None:hist.append(["Tarih UTC","Takım","İşlem","Oyuncu","Oyuncu ID"])
 for event in changes:hist.append(list(event))
 meta=wb["SENKRON_DURUMU"] if "SENKRON_DURUMU" in wb else wb.create_sheet("SENKRON_DURUMU")
 meta["A1"]="Son başarılı kontrol (UTC)"
 meta["B1"]=now
 meta["A2"]="Lig ID"
 meta["B2"]=88948640
 meta["A3"]="Takım sayısı"
 meta["B3"]=len(teams)
 wb.save(book)
 print("Synced",len(teams),"teams",len(changes),"roster changes")
if __name__=="__main__":main()
